"""Recibe archivos sueltos de la raíz; valida, conserva versiones y organiza."""

from runtime_paths import DATA_DIR, CODE_DIR, APP_DIR, APP_URL, company_rut, control_value, category_provider_rut
from datetime import datetime
import hashlib
import json
from pathlib import Path
import shutil

from openpyxl import load_workbook
from conciliar import BASE, read_excel, normal
from xml_source import read_xml_documents, read_zip
from zeta_bot import exclusive_lock


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def classify(path):
    suffix = path.suffix.lower()
    if suffix == '.json':
        data = json.loads(path.read_text(encoding='utf-8-sig'))
        if not isinstance(data,dict):
            raise ValueError('JSON no reconocido: ' + path.name)
        if data.get('formato')=='cerecito-catalogos-v1' and data.get('version')==1 and isinstance(data.get('catalogs'),dict):
            return 'catalogos'
        if data.get('formato') == 'cerecito-agente-v1' and isinstance(data.get('comprobantes'), list):
            keys = [d.get('clave') for d in data['comprobantes']]
            if not keys or any(not k for k in keys) or len(set(keys)) != len(keys):
                raise ValueError('JSON vacío o con identidades repetidas: ' + path.name)
            return 'cerecito'
        if data.get('version') == 3 and isinstance(data.get('manualRules'), dict) and isinstance(data.get('providerOverrides'), dict):
            return 'correlaciones'
        raise ValueError('JSON no reconocido: ' + path.name)
    if suffix == '.zip':
        if not read_zip(path):
            raise ValueError('ZIP sin XML de facturas: ' + path.name)
        return 'xml'
    if suffix == '.xml':
        if path.stat().st_size > 8 * 1024 * 1024:
            raise ValueError('XML supera el tamaño permitido.')
        read_xml_documents(path.read_bytes(), path.name)
        return 'xml'
    if suffix == '.xlsx':
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            header = [str(x or '') for x in next(book.active.iter_rows(min_row=2, max_row=2, values_only=True))]
            if 'Tipo (C|L)' in header or 'Tipo(C|L)' in header:
                return 'referencia_articulos'
            for sheet in book.worksheets:
                for row in sheet.iter_rows(min_row=1,max_row=min(sheet.max_row,20),values_only=True):
                    fields={normal(value) for value in row if value is not None}
                    if ('CODIGO' in fields and fields & {'NOMBRE','DESCRIPCION'}
                            and not fields & {'PROVEEDOR','TOTAL','Nº','NRO','NUMERO'}):
                        return 'catalogos'
        finally:
            book.close()
        rows, _ = read_excel(path)
        if not rows:
            raise ValueError('Excel sin compras reconocidas: ' + path.name)
        return 'zeta'
    return None


def archive(path, output, group):
    folder = output / 'entradas_anteriores' / group
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / (path.stem + '_' + sha(path)[:16] + path.suffix)
    if target.exists():
        if sha(target) != sha(path):
            raise ValueError('Conflicto de respaldo: ' + path.name)
        target = folder / (path.stem + '_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f') + path.suffix)
    shutil.move(str(path), str(target))
    return target


def activate(source, inputs, output, group, *, copy=False):
    folder = inputs / group
    folder.mkdir(parents=True, exist_ok=True)
    digest = sha(source)
    matches = [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() == source.suffix.lower()]
    # JSON y Excel histórico: una sola versión activa. Versiones previas se conservan.
    if group in {'cerecito', 'zeta', 'correlaciones'}:
        for old in matches:
            if sha(old) == digest:
                if not copy and old.resolve() != source.resolve():
                    archive(source, output, 'duplicados_recibidos')
                return old
        for old in matches:
            archive(old, output, group)
    target = folder / source.name
    if target.exists():
        if sha(target) == digest:
            if not copy and target.resolve() != source.resolve():
                archive(source, output, 'duplicados_recibidos')
            return target
        archive(target, output, group)
    if copy:
        temp = target.with_suffix(target.suffix + '.tmp')
        if temp.exists():
            raise ValueError('Hay una copia interrumpida pendiente de revisión: ' + temp.name)
        shutil.copy2(source, temp)
        if sha(temp) != digest:
            raise ValueError('No coincide la copia de entrada.')
        temp.replace(target)
    else:
        shutil.move(str(source), str(target))
    if sha(target) != digest:
        raise ValueError('Cambió el archivo durante la recepción: ' + target.name)
    return target


def receive(root):
    inputs, output = DATA_DIR / 'entrada', DATA_DIR / 'salida'
    plan = []
    for path in sorted(root.iterdir()):
        if not path.is_file() or path.suffix.lower() not in {'.json', '.zip', '.xml', '.xlsx'}:
            continue
        group = classify(path)
        plan.append({'origen': str(path), 'grupo': group, 'sha256': sha(path)})
    for group in ('cerecito', 'zeta', 'correlaciones'):
        digests = {p['sha256'] for p in plan if p['grupo'] == group}
        if len(digests) > 1:
            raise ValueError('Hay varias versiones nuevas de ' + group + ' en la raíz; no se elige automáticamente.')
    folder = output / 'recepcion'
    folder.mkdir(parents=True, exist_ok=True)
    # Una recepción incompleta no se repite silenciosamente.
    for journal in folder.glob('*.json'):
        if json.loads(journal.read_text(encoding='utf-8')).get('estado') == 'EN_CURSO':
            raise ValueError('Hay una recepción interrumpida: revisar ' + str(journal))
    journal = folder / ('recepcion_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f') + '.json')
    report = {'estado': 'EN_CURSO', 'archivos': plan}
    journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    for item in plan:
        source = Path(item['origen'])
        if sha(source) != item['sha256']:
            raise ValueError('La entrada cambió durante la recepción.')
        destination = activate(source, inputs, output, item['grupo'])
        item['destino'] = str(destination)
        journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print('Recibido: ' + source.name + ' -> archivos/' + item['grupo'], flush=True)
    report['estado'] = 'RECIBIDO_NO_EQUIVALE_A_CARGADO_EN_ZETA'
    journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    (folder / 'ultima.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report


if __name__ == '__main__':
    try:
        with exclusive_lock(BASE / 'estado/recepcion.lock'):
            receive(DATA_DIR/'importar')
    except Exception as exc:
        print('RECEPCIÓN BLOQUEADA:', str(exc))
        raise SystemExit(1)

